from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


OUTPUT_FILE = "SRE_Operational_Data.xlsx"


def style_sheet(ws):
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for cell in ws[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center"
        )

    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value)),
                )

        ws.column_dimensions[
            column_letter
        ].width = min(max_length + 3, 40)


def create_workbook():
    workbook = Workbook()

    # ---------------------------------------------------------
    # README
    # ---------------------------------------------------------

    ws = workbook.active
    ws.title = "README"

    ws.append(
        [
            "Field",
            "Value",
        ]
    )

    ws.append(
        [
            "Workbook",
            "SRE Operational Data",
        ]
    )

    ws.append(
        [
            "Purpose",
            "Enterprise SRE incident investigation and operational analysis",
        ]
    )

    ws.append(
        [
            "Primary Incident",
            "INC-1003",
        ]
    )

    ws.append(
        [
            "Important Scenario",
            "Payment API degradation caused by database connection saturation",
        ]
    )

    ws.append(
        [
            "Data Sources",
            "Incidents, Service Metrics, Dependencies, Historical RCA, Infrastructure",
        ]
    )

    style_sheet(ws)

    # ---------------------------------------------------------
    # INCIDENTS
    # ---------------------------------------------------------

    ws = workbook.create_sheet("Incidents")

    ws.append(
        [
            "Incident_ID",
            "Timestamp",
            "Service",
            "Severity",
            "Status",
            "Error_Rate",
            "Latency_ms",
            "Affected_Users",
            "Root_Cause",
        ]
    )

    incidents = [
        [
            "INC-1001",
            "2026-08-25 09:15",
            "Authentication Service",
            "P2",
            "Resolved",
            4.2,
            620,
            1200,
            "Redis connection exhaustion",
        ],
        [
            "INC-1002",
            "2026-08-27 14:30",
            "Order Service",
            "P2",
            "Resolved",
            3.8,
            710,
            850,
            "Downstream inventory timeout",
        ],
        [
            "INC-1003",
            "2026-08-30 16:45",
            "Payment API",
            "P1",
            "Investigating",
            12.7,
            1850,
            5400,
            None,
        ],
        [
            "INC-1004",
            "2026-09-01 11:20",
            "Notification Service",
            "P3",
            "Resolved",
            2.1,
            480,
            500,
            "SMTP provider timeout",
        ],
    ]

    for row in incidents:
        ws.append(row)

    style_sheet(ws)

    # ---------------------------------------------------------
    # SERVICE METRICS
    # ---------------------------------------------------------

    ws = workbook.create_sheet(
        "Service_Metrics"
    )

    ws.append(
        [
            "Timestamp",
            "Service",
            "CPU_Percent",
            "Memory_Percent",
            "DB_Connection_Utilization",
            "Request_Rate",
            "Error_Rate",
            "Latency_ms",
        ]
    )

    metrics = [
        [
            "2026-08-30 16:15",
            "Payment API",
            58,
            64,
            61,
            820,
            1.2,
            420,
        ],
        [
            "2026-08-30 16:30",
            "Payment API",
            62,
            67,
            74,
            850,
            3.5,
            690,
        ],
        [
            "2026-08-30 16:45",
            "Payment API",
            71,
            72,
            91,
            910,
            12.7,
            1850,
        ],
        [
            "2026-08-30 17:00",
            "Payment API",
            73,
            74,
            95,
            940,
            15.2,
            2210,
        ],
        [
            "2026-08-30 17:15",
            "Payment API",
            69,
            71,
            93,
            900,
            11.4,
            1760,
        ],
        [
            "2026-08-30 16:45",
            "Database",
            84,
            78,
            96,
            1250,
            8.9,
            1420,
        ],
        [
            "2026-08-30 17:00",
            "Database",
            89,
            81,
            98,
            1310,
            11.8,
            1690,
        ],
        [
            "2026-08-30 17:15",
            "Database",
            86,
            79,
            94,
            1270,
            9.7,
            1510,
        ],
        [
            "2026-08-30 16:45",
            "API Gateway",
            48,
            55,
            42,
            930,
            2.1,
            520,
        ],
    ]

    for row in metrics:
        ws.append(row)

    style_sheet(ws)

    # ---------------------------------------------------------
    # DEPENDENCIES
    # ---------------------------------------------------------

    ws = workbook.create_sheet(
        "Dependencies"
    )

    ws.append(
        [
            "Source_Service",
            "Target_Service",
            "Dependency_Type",
            "Criticality",
            "Relationship",
        ]
    )

    dependencies = [
        [
            "Payment API",
            "API Gateway",
            "Inbound",
            "High",
            "Payment API receives traffic through API Gateway",
        ],
        [
            "Payment API",
            "Authentication Service",
            "Synchronous",
            "High",
            "Payment API validates authenticated requests",
        ],
        [
            "Payment API",
            "Database",
            "Synchronous",
            "Critical",
            "Payment transactions require database access",
        ],
        [
            "Database",
            "DB Connection Pool",
            "Internal",
            "Critical",
            "Database relies on connection pool",
        ],
        [
            "Database",
            "DB Replica",
            "Replication",
            "High",
            "Primary database replicates transaction data",
        ],
        [
            "Order Service",
            "Payment API",
            "Synchronous",
            "High",
            "Order processing depends on payment confirmation",
        ],
    ]

    for row in dependencies:
        ws.append(row)

    style_sheet(ws)

    # ---------------------------------------------------------
    # HISTORICAL RCA
    # ---------------------------------------------------------

    ws = workbook.create_sheet(
        "Historical_RCA"
    )

    ws.append(
        [
            "Incident_ID",
            "Service",
            "Symptom",
            "Root_Cause",
            "Resolution",
            "Preventive_Action",
        ]
    )

    historical_rca = [
        [
            "INC-0901",
            "Payment API",
            "High payment latency",
            "Database connection pool exhaustion",
            "Increased connection pool capacity and restarted affected workers",
            "Connection pool monitoring",
        ],
        [
            "INC-0912",
            "Payment API",
            "Intermittent transaction failures",
            "Database saturation during peak traffic",
            "Optimized database queries and scaled database resources",
            "Database utilization alerting",
        ],
        [
            "INC-0955",
            "Database",
            "High query latency",
            "Excessive concurrent connections",
            "Connection throttling and query optimization",
            "Connection utilization dashboard",
        ],
        [
            "INC-0988",
            "Order Service",
            "Transaction timeout",
            "Payment API dependency latency",
            "Payment API timeout tuning",
            "Dependency latency monitoring",
        ],
    ]

    for row in historical_rca:
        ws.append(row)

    style_sheet(ws)

    # ---------------------------------------------------------
    # INFRASTRUCTURE
    # ---------------------------------------------------------

    ws = workbook.create_sheet(
        "Infrastructure"
    )

    ws.append(
        [
            "Component_ID",
            "Component_Name",
            "Component_Type",
            "Environment",
            "Region",
            "Capacity",
            "Current_Utilization",
            "Health_Status",
        ]
    )

    infrastructure = [
        [
            "INF-001",
            "payment-api-prod",
            "Application",
            "Production",
            "ap-south-1",
            "1000 req/s",
            "91%",
            "Degraded",
        ],
        [
            "INF-002",
            "payment-db-primary",
            "Database",
            "Production",
            "ap-south-1",
            "500 connections",
            "98%",
            "Critical",
        ],
        [
            "INF-003",
            "payment-db-replica",
            "Database Replica",
            "Production",
            "ap-south-1",
            "500 connections",
            "64%",
            "Healthy",
        ],
        [
            "INF-004",
            "api-gateway-prod",
            "API Gateway",
            "Production",
            "ap-south-1",
            "2000 req/s",
            "47%",
            "Healthy",
        ],
        [
            "INF-005",
            "auth-service-prod",
            "Application",
            "Production",
            "ap-south-1",
            "1200 req/s",
            "52%",
            "Healthy",
        ],
    ]

    for row in infrastructure:
        ws.append(row)

    style_sheet(ws)

    # ---------------------------------------------------------
    # SAVE
    # ---------------------------------------------------------

    workbook.save(OUTPUT_FILE)

    print(
        f"Excel workbook created successfully: {OUTPUT_FILE}"
    )


if __name__ == "__main__":
    create_workbook()